from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = r"D:\DA26TTNT\DA26TTNT-NMTTNT\dự án\Tai_lieu_giai_thich_ma_nguon_DA26TTNT.docx"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F6F7"
GRID = "D9D9D9"
TEXT = RGBColor(31, 41, 55)


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), GRID)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    run.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    set_table_borders(table)
    for i, heading in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        r = p.add_run(heading)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_fill(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(9.2)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    p.add_run(text)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.7)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12

for style_name, size, before, after in (("Title", 26, 0, 16), ("Heading 1", 18, 16, 8), ("Heading 2", 13, 11, 5), ("Heading 3", 11, 8, 4)):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name in ("Title", "Heading 1") else "Aptos"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

if "Code Block" not in [s.name for s in styles]:
    code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
else:
    code_style = styles["Code Block"]
code_style.font.name = "Consolas"
code_style.font.size = Pt(8.6)
code_style.font.color.rgb = RGBColor(30, 41, 59)
code_style.paragraph_format.left_indent = Inches(0.25)
code_style.paragraph_format.right_indent = Inches(0.18)
code_style.paragraph_format.space_before = Pt(4)
code_style.paragraph_format.space_after = Pt(8)
code_style.paragraph_format.line_spacing = 1.0
code_ppr = code_style.element.get_or_add_pPr()
shd = OxmlElement("w:shd")
shd.set(qn("w:fill"), PALE_GRAY)
code_ppr.append(shd)

footer = section.footer.paragraphs[0]
footer.add_run("DA26TTNT   Hướng dẫn mã nguồn cho người mới")
footer.runs[0].font.size = Pt(8.5)
footer.runs[0].font.color.rgb = RGBColor(100, 116, 139)
add_page_number(footer)

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Hướng dẫn mã nguồn website DA26TTNT")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Giải thích kiến trúc và cách hoạt động dành cho người mới")
r.font.size = Pt(15)
r.bold = True
p.paragraph_format.space_after = Pt(30)

add_table(doc, ["Thông tin", "Nội dung"], [
    ["Loại tài liệu", "Sổ tay kỹ thuật và hướng dẫn vận hành"],
    ["Đối tượng", "Người mới học Next.js, TypeScript, Prisma và PostgreSQL"],
    ["Dự án", "DA26TTNT Cổng thông tin lớp"],
    ["Phiên bản mã nguồn", "Trạng thái tại ngày 08 tháng 10 năm 2026"],
], [1.6, 5.0])

doc.add_heading("Mục tiêu của tài liệu", level=1)
doc.add_paragraph("Tài liệu này giúp bạn hiểu website đang gồm những phần nào, một yêu cầu đi qua hệ thống ra sao, dữ liệu được lưu ở đâu và cách chạy dự án trên máy Windows. Sau khi đọc, bạn có thể tìm đúng file để sửa giao diện, thay đổi cơ sở dữ liệu, điều chỉnh đăng nhập và tiếp tục xây dựng các chức năng còn thiếu.")
doc.add_paragraph("Mã nguồn hiện đã có giao diện, mô hình dữ liệu, đăng nhập bằng cơ sở dữ liệu, session cookie, bảo vệ tuyến quản trị, migration đầu tiên và tài khoản quản trị local. Các màn hình quản trị chuyên sâu, CRUD đầy đủ và quy trình nhập Excel hoặc CSV mới có khung giao diện, chưa phải chức năng hoàn chỉnh.")

doc.add_page_break()

doc.add_heading("Mục lục", level=1)
toc_items = [
    "1 Tổng quan website", "2 Mô hình hoạt động", "3 Công nghệ sử dụng", "4 Cấu trúc thư mục",
    "5 Giao diện và định tuyến", "6 Đăng nhập và phiên làm việc", "7 Phân quyền và middleware",
    "8 Cơ sở dữ liệu Prisma", "9 Khởi tạo dữ liệu và migration", "10 Cách chạy website",
    "11 Cách đọc và sửa mã nguồn", "12 Kiểm tra và xử lý lỗi", "13 Giới hạn hiện tại",
    "14 Từ điển thuật ngữ", "15 Bảng tra cứu từng file"
]
for item in toc_items:
    p = doc.add_paragraph(item)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(4)

doc.add_heading("1 Tổng quan website", level=1)
doc.add_paragraph("DA26TTNT là website cổng thông tin lớp chạy trên trình duyệt. Next.js đảm nhiệm cả giao diện và phần xử lý phía máy chủ. PostgreSQL giữ dữ liệu lâu dài. Prisma là lớp trung gian giúp mã TypeScript đọc và ghi PostgreSQL mà không phải viết SQL cho mọi thao tác.")
add_table(doc, ["Phần", "Vai trò", "Ví dụ trong dự án"], [
    ["Trình duyệt", "Hiển thị giao diện và gửi thao tác của người dùng", "Trang chủ, thành viên, thông báo, đăng nhập"],
    ["Next.js", "Kết xuất trang và chạy mã máy chủ", "app/page.tsx và Server Action đăng nhập"],
    ["PostgreSQL", "Lưu tài khoản, phiên, thành viên và thông báo", "Database da26ttnt"],
    ["Prisma", "Ánh xạ bảng thành đối tượng TypeScript", "prisma/schema.prisma và lib/prisma.ts"],
], [1.2, 2.6, 2.8])

doc.add_heading("2 Mô hình hoạt động", level=1)
doc.add_heading("2.1 Luồng của một yêu cầu", level=2)
add_code(doc, "Người dùng → Trình duyệt → Middleware → Trang hoặc Server Action → Prisma → PostgreSQL\n                 ← HTML hoặc kết quả thao tác ←              ← dữ liệu ←")
doc.add_paragraph("Khi người dùng mở một URL, Next.js nhận yêu cầu. Middleware kiểm tra nhanh xem URL có thuộc nhóm cần đăng nhập hay không. Trang tĩnh có thể được trả về ngay. Trang hoặc thao tác cần dữ liệu gọi Prisma, Prisma gửi câu lệnh đến PostgreSQL rồi trả kết quả về mã TypeScript.")

doc.add_heading("2.2 Ranh giới phía trình duyệt và phía máy chủ", level=2)
add_table(doc, ["Nơi chạy", "Dấu hiệu", "Được làm", "Không nên làm"], [
    ["Trình duyệt", "File có dòng use client", "Xử lý nút bấm, trạng thái form, hiệu ứng", "Giữ mật khẩu database hoặc tự quyết định quyền"],
    ["Máy chủ", "Server Component hoặc use server", "Truy vấn DB, kiểm tra mật khẩu, tạo session", "Gửi passwordHash hoặc bí mật ra HTML"],
], [1.1, 1.5, 2.0, 2.0])

doc.add_heading("3 Công nghệ sử dụng", level=1)
add_table(doc, ["Công nghệ", "Mục đích", "Giải thích ngắn"], [
    ["Next.js 15", "Khung ứng dụng web", "App Router biến thư mục app thành URL và hỗ trợ mã máy chủ"],
    ["React 19", "Xây dựng giao diện", "Mỗi phần giao diện là một component có thể tái sử dụng"],
    ["TypeScript", "Kiểm tra kiểu", "Phát hiện sai kiểu dữ liệu trước khi chạy"],
    ["Prisma 6", "Truy cập dữ liệu", "Sinh Prisma Client từ schema và quản lý migration"],
    ["PostgreSQL 17", "Cơ sở dữ liệu", "Lưu dữ liệu quan hệ bền vững"],
    ["bcryptjs", "Bảo vệ mật khẩu", "So sánh mật khẩu với giá trị băm"],
    ["Zod", "Kiểm tra đầu vào", "Xác nhận email và độ dài mật khẩu ở máy chủ"],
    ["Lucide React", "Biểu tượng", "Cung cấp icon nhất quán cho giao diện"],
    ["XLSX", "Đọc bảng tính", "Đã cài dependency nhưng luồng nhập file chưa hoàn chỉnh"],
], [1.35, 1.55, 3.65])

doc.add_page_break()
doc.add_heading("4 Cấu trúc thư mục", level=1)
add_code(doc, "dự án/\n├─ app/                 Trang và Server Action theo App Router\n├─ components/          Component giao diện dùng chung\n├─ lib/                 Kết nối Prisma, session và hàm tiện ích\n├─ prisma/              Schema, migration và script seed\n├─ storage/             Điểm dành cho file local trong tương lai\n├─ tools/               Công cụ tạo tài liệu kỹ thuật\n├─ middleware.ts        Chặn URL cần đăng nhập\n├─ package.json         Dependency và lệnh npm\n├─ next.config.ts       Cấu hình Next.js\n├─ tsconfig.json        Cấu hình TypeScript\n├─ .env                 Bí mật local, không đưa lên Git\n└─ SETUP.md             Hướng dẫn cài đặt ngắn")
doc.add_paragraph("Trong App Router, tên thư mục trở thành một phần của URL. Ví dụ app/thanh-vien/page.tsx tạo URL /thanh-vien. File layout.tsx là khung HTML gốc. File globals.css chứa kiểu dáng dùng cho toàn website.")

doc.add_heading("5 Giao diện và định tuyến", level=1)
doc.add_heading("5.1 Các trang hiện có", level=2)
add_table(doc, ["URL", "File", "Mục đích"], [
    ["/", "app/page.tsx", "Trang chủ và số liệu tổng quan"],
    ["/dang-nhap", "app/dang-nhap/page.tsx", "Màn hình đăng nhập"],
    ["/thanh-vien", "app/thanh-vien/page.tsx", "Danh sách và bộ lọc thành viên"],
    ["/thong-bao", "app/thong-bao/page.tsx", "Danh sách thông báo"],
    ["/co-van", "app/co-van/page.tsx", "Thông tin cố vấn học tập"],
    ["/ban-can-su", "app/ban-can-su/page.tsx", "Danh sách ban cán sự"],
    ["/lich", "app/lich/page.tsx", "Lịch lớp"],
    ["/tai-lieu", "app/tai-lieu/page.tsx", "Tài liệu dùng chung"],
    ["/quan-tri", "app/quan-tri/page.tsx", "Tổng quan quản trị"],
    ["/quan-tri/[section]", "app/quan-tri/[section]/page.tsx", "Khung cho từng khu vực quản trị"],
], [1.4, 2.35, 2.8])

doc.add_heading("5.2 Component dùng chung", level=2)
doc.add_paragraph("AppShell tạo thanh điều hướng bên trái, thanh tìm kiếm phía trên và vùng nội dung. PageHead chuẩn hóa tiêu đề trang. EmptyState cung cấp trạng thái chưa có dữ liệu. Việc tách component giúp sửa một lần và áp dụng cho nhiều trang.")

doc.add_heading("5.3 CSS và responsive", level=2)
doc.add_paragraph("app/globals.css định nghĩa màu, khoảng cách, card, bảng, form và breakpoint. Ở màn hình dưới 900 px, sidebar thu nhỏ còn icon. Dưới 620 px, thanh điều hướng chuyển xuống cạnh dưới để phù hợp điện thoại. Đây là CSS thuần, chưa dùng Tailwind dù yêu cầu ban đầu có nhắc Tailwind CSS.")

doc.add_heading("6 Đăng nhập và phiên làm việc", level=1)
doc.add_heading("6.1 Luồng đăng nhập", level=2)
add_code(doc, "1  Người dùng gửi email và mật khẩu\n2  Zod kiểm tra định dạng\n3  Prisma tìm User theo email\n4  bcryptjs so sánh mật khẩu với passwordHash\n5  Hệ thống tạo token ngẫu nhiên 32 byte\n6  Chỉ mã băm SHA 256 của token được lưu trong Session\n7  Token gốc được đặt trong cookie httpOnly trong 7 ngày\n8  Trình duyệt tự gửi cookie ở yêu cầu tiếp theo")
doc.add_paragraph("Form đăng nhập là Client Component vì nó dùng useActionState để hiển thị trạng thái đang gửi và lỗi. Hàm login trong app/dang-nhap/actions.ts là Server Action nên mật khẩu được xử lý ở máy chủ. Mật khẩu trong database là bcrypt hash, không phải văn bản gốc.")

doc.add_heading("6.2 Các thuộc tính cookie", level=2)
add_table(doc, ["Thuộc tính", "Ý nghĩa"], [
    ["httpOnly", "JavaScript trong trình duyệt không đọc được token"],
    ["sameSite lax", "Giảm nguy cơ gửi cookie trong một số yêu cầu từ website khác"],
    ["secure", "Chỉ gửi qua HTTPS khi chạy production"],
    ["expires", "Phiên hiện có thời hạn bảy ngày"],
], [1.5, 5.1])

doc.add_heading("7 Phân quyền và middleware", level=1)
doc.add_paragraph("middleware.ts kiểm tra sự tồn tại của cookie trước khi cho truy cập các URL nội bộ. Đây là lớp chặn sớm để chuyển người chưa đăng nhập về /dang-nhap. Cookie tồn tại chưa chứng minh nó hợp lệ, vì vậy mã máy chủ vẫn phải gọi currentUser hoặc requireAdmin trước thao tác nhạy cảm.")
add_table(doc, ["Vai trò", "Ý nghĩa trong schema", "Tình trạng hiện tại"], [
    ["MEMBER", "Thành viên thông thường", "Mô hình đã có, chức năng cập nhật hồ sơ chưa hoàn chỉnh"],
    ["CLASS_OFFICER", "Ban cán sự", "Mô hình đã có, phạm vi quyền chi tiết chưa hoàn chỉnh"],
    ["ADMIN", "Quản trị viên", "Trang /quan-tri gọi requireAdmin"],
], [1.35, 2.05, 3.15])
doc.add_paragraph("Điểm cần sửa tiếp: trang quản trị động app/quan-tri/[section]/page.tsx hiện dựa vào middleware nhưng chưa tự gọi requireAdmin. Mọi Server Action hoặc API mới phải kiểm tra quyền độc lập, không được dựa vào việc ẩn nút trên giao diện.")

doc.add_page_break()
doc.add_heading("8 Cơ sở dữ liệu Prisma", level=1)
doc.add_heading("8.1 Cách đọc schema", level=2)
doc.add_paragraph("Mỗi khối model trong prisma/schema.prisma tương ứng với một bảng. @id là khóa chính, @unique cấm trùng, @default đặt giá trị mặc định, @relation mô tả quan hệ và @@index tạo chỉ mục để truy vấn nhanh hơn.")
add_table(doc, ["Model", "Dữ liệu giữ", "Quan hệ chính"], [
    ["User", "Email, passwordHash, vai trò, trạng thái", "Có thể liên kết một Member, nhiều Session và Announcement"],
    ["Session", "Mã băm token và thời hạn", "Thuộc về một User"],
    ["Member", "Mã sinh viên, họ tên và thông tin công khai", "Có thể liên kết User và nhiều ClassPosition"],
    ["ClassPosition", "Chức vụ và nhiệm vụ", "Thuộc về một Member"],
    ["Advisor", "Thông tin cố vấn và lịch tư vấn", "Độc lập trong schema hiện tại"],
    ["Announcement", "Nội dung, trạng thái, độ quan trọng", "Có tác giả, file và danh sách đã đọc"],
    ["Attachment", "Metadata của file", "Thuộc về một Announcement"],
    ["AnnouncementRead", "Ai đã đọc và lúc nào", "Bảng nối User và Announcement"],
    ["ImportJob", "Kết quả mỗi lần nhập dữ liệu", "Hiện chỉ giữ createdById dạng chuỗi"],
    ["AuditLog", "Hành động quản trị quan trọng", "Có thể liên kết User thực hiện"],
], [1.3, 2.25, 3.05])

doc.add_heading("8.2 Quan hệ quan trọng", level=2)
add_code(doc, "User 1 ─ 0..1 Member\nUser 1 ─ nhiều Session\nMember 1 ─ nhiều ClassPosition\nUser 1 ─ nhiều Announcement\nAnnouncement 1 ─ nhiều Attachment\nUser nhiều ─ nhiều Announcement qua AnnouncementRead")
doc.add_paragraph("onDelete Cascade xóa bản ghi con khi bản ghi cha bị xóa. onDelete SetNull giữ bản ghi nhưng bỏ liên kết. Hai lựa chọn này ảnh hưởng trực tiếp đến an toàn dữ liệu, vì vậy cần xem migration trước khi áp dụng trên production.")

doc.add_heading("9 Khởi tạo dữ liệu và migration", level=1)
doc.add_heading("9.1 Migration", level=2)
doc.add_paragraph("Migration là lịch sử thay đổi cấu trúc database. Migration đầu tiên nằm trong prisma/migrations/20261007182302_init/migration.sql. Lệnh migrate dev so sánh schema, dùng shadow database để kiểm tra, tạo file SQL rồi áp dụng vào database local.")
add_code(doc, "npm.cmd run db:generate\nnpm.cmd run db:migrate -- --name ten_thay_doi")
doc.add_paragraph("Không sửa migration đã chạy trên môi trường dùng chung. Khi cần đổi model, sửa schema.prisma rồi tạo migration mới với tên mô tả ngắn.")

doc.add_heading("9.2 Seed quản trị viên", level=2)
doc.add_paragraph("prisma/seed.ts đọc BOOTSTRAP_ADMIN_EMAIL và BOOTSTRAP_ADMIN_PASSWORD. Script chỉ chạy khi có email và mật khẩu dài ít nhất 12 ký tự. Mật khẩu được băm bằng bcrypt trước khi lưu. upsert giúp chạy lại mà không tạo tài khoản trùng.")

doc.add_heading("9.3 Biến môi trường", level=2)
add_code(doc, "DATABASE_URL=postgresql://USER:PASSWORD@127.0.0.1:5432/da26ttnt?schema=public\nSHADOW_DATABASE_URL=postgresql://USER:PASSWORD@127.0.0.1:5432/da26ttnt_shadow?schema=public\nSESSION_SECRET=mot_chuoi_ngau_nhien_dai\nBOOTSTRAP_ADMIN_EMAIL=dia_chi_quan_tri\nBOOTSTRAP_ADMIN_PASSWORD=mat_khau_khoi_tao")
doc.add_paragraph(".env chứa bí mật local và đã nằm trong .gitignore. Không chép mật khẩu thật vào tài liệu, commit Git, ảnh chụp màn hình hoặc log. .env.example chỉ dùng giá trị mẫu để người khác biết cần cấu hình biến nào.")

doc.add_page_break()
doc.add_heading("10 Cách chạy website", level=1)
doc.add_heading("10.1 Chạy để phát triển", level=2)
add_code(doc, "cd \"D:\\DA26TTNT\\DA26TTNT-NMTTNT\\dự án\"\n$env:Path = \"C:\\Program Files\\nodejs;$env:Path\"\nnpm.cmd install\nnpm.cmd run db:generate\nnpm.cmd run dev")
doc.add_paragraph("Mở http://localhost:3000. Giữ cửa sổ PowerShell đang chạy. Nhấn Ctrl C để dừng server. PostgreSQL được cài thành Windows Service nên thường tự khởi động cùng máy.")

doc.add_heading("10.2 Chạy bản production local", level=2)
add_code(doc, "npm.cmd run build\nnpm.cmd run start")
doc.add_paragraph("Lệnh build biên dịch và kiểm tra dự án. Lệnh start chạy kết quả đã build. Nếu sửa code, phải build lại trước khi chạy bản production mới.")

doc.add_heading("10.3 Các lệnh thường dùng", level=2)
add_table(doc, ["Lệnh", "Khi nào dùng"], [
    ["npm.cmd run dev", "Phát triển với tự tải lại khi lưu file"],
    ["npm.cmd run lint", "Kiểm tra quy tắc mã nguồn"],
    ["npm.cmd run typecheck", "Kiểm tra kiểu TypeScript"],
    ["npm.cmd run build", "Xác minh bản production"],
    ["npm.cmd run db:generate", "Sinh lại Prisma Client sau khi đổi schema"],
    ["npm.cmd run db:migrate -- --name ten", "Tạo và áp migration mới"],
    ["npm.cmd run db:seed", "Tạo hoặc bảo đảm tài khoản quản trị đầu tiên"],
], [2.8, 3.7])

doc.add_heading("11 Cách đọc và sửa mã nguồn", level=1)
doc.add_heading("11.1 Sửa nội dung một trang", level=2)
add_bullets(doc, [
    "Tìm URL cần sửa và mở file page.tsx tương ứng trong app.",
    "Sửa JSX để thay nội dung hoặc cấu trúc hiển thị.",
    "Nếu nhiều trang cùng dùng một phần, chuyển phần đó thành component trong components.",
    "Sửa app/globals.css khi cần đổi kiểu dáng chung hoặc responsive.",
    "Chạy lint, typecheck và build sau thay đổi đáng kể.",
])

doc.add_heading("11.2 Thêm trường dữ liệu", level=2)
add_bullets(doc, [
    "Thêm trường vào model phù hợp trong prisma/schema.prisma.",
    "Tạo migration mới thay vì sửa database thủ công.",
    "Chạy db:generate để TypeScript biết trường mới.",
    "Cập nhật Zod schema để kiểm tra đầu vào.",
    "Cập nhật Server Action hoặc API để đọc và ghi trường.",
    "Cập nhật trang hiển thị và form nhập liệu.",
])

doc.add_heading("11.3 Ví dụ thêm thông báo", level=2)
doc.add_paragraph("Một chức năng tạo thông báo hoàn chỉnh nên có form phía trình duyệt, Zod schema phía máy chủ, Server Action kiểm tra currentUser và vai trò, lệnh prisma.announcement.create, ghi AuditLog và thông báo kết quả. Giao diện chỉ báo thành công sau khi transaction hoàn tất.")

doc.add_heading("12 Kiểm tra và xử lý lỗi", level=1)
add_table(doc, ["Hiện tượng", "Nguyên nhân thường gặp", "Cách kiểm tra"], [
    ["node hoặc npm không nhận diện", "Terminal chưa nạp PATH", "Thêm C:\\Program Files\\nodejs vào PATH của phiên"],
    ["npm.ps1 bị chặn", "PowerShell ExecutionPolicy", "Dùng npm.cmd thay cho npm"],
    ["P1001 hoặc không kết nối DB", "PostgreSQL dừng hoặc URL sai", "Kiểm tra service và DATABASE_URL"],
    ["P3014 shadow database", "Thiếu SHADOW_DATABASE_URL hoặc quyền", "Kiểm tra database da26ttnt_shadow"],
    ["Đăng nhập thất bại", "Sai email, mật khẩu, trạng thái hoặc hash", "Kiểm tra User có status ACTIVE"],
    ["Vào trang rồi bị chuyển đăng nhập", "Không có cookie session", "Đăng nhập lại và kiểm tra cookie"],
    ["Build lỗi", "Sai kiểu hoặc import", "Chạy typecheck và đọc dòng lỗi đầu tiên"],
], [1.65, 2.2, 2.65])

doc.add_heading("12.1 Bộ kiểm tra tối thiểu", level=2)
add_code(doc, "npm.cmd run lint\nnpm.cmd run typecheck\nnpm.cmd run build\nnpm.cmd exec prisma migrate status")
doc.add_paragraph("Tại thời điểm viết tài liệu, Prisma generate, lint, typecheck và production build đã chạy thành công. Migration đầu tiên đã được áp vào PostgreSQL local.")

doc.add_page_break()
doc.add_heading("13 Giới hạn hiện tại", level=1)
doc.add_paragraph("Phần này mô tả trạng thái mã nguồn thực tế để người mới không nhầm giao diện mẫu với chức năng đã hoàn thiện.")
add_table(doc, ["Hạng mục", "Trạng thái", "Việc cần làm"], [
    ["Trang giao diện chính", "Đã có", "Kết nối số liệu thật từ Prisma"],
    ["Đăng nhập quản trị", "Đã có nền tảng", "Thêm đăng xuất, đổi mật khẩu, giới hạn tần suất"],
    ["Phân quyền", "Có role và requireAdmin", "Áp kiểm tra cho mọi action và trang động quản trị"],
    ["Quản lý thành viên", "Mới có màn hình", "Xây CRUD, phân trang, tìm không dấu"],
    ["Thông báo", "Mới có model và màn hình", "Xây soạn thảo, trạng thái, ghim, đã đọc"],
    ["Nhập Excel hoặc CSV", "Mới có dependency và khung", "Xây upload, mapping, preview, transaction và báo lỗi"],
    ["File đính kèm", "Mới có model", "Xây storage adapter, kiểm tra MIME, kích thước và quyền tải"],
    ["Email mời", "Chưa có", "Thêm adapter email và quy trình kích hoạt"],
    ["Tailwind CSS", "Chưa dùng", "Giữ CSS hiện tại hoặc chuyển đổi có kế hoạch"],
    ["Kiểm thử tự động", "Chưa có test suite", "Bổ sung unit, integration và end to end"],
], [1.65, 1.45, 3.4])

doc.add_heading("14 Từ điển thuật ngữ", level=1)
add_table(doc, ["Thuật ngữ", "Giải thích cho người mới"], [
    ["Route", "Đường dẫn URL dẫn đến một trang hoặc endpoint"],
    ["Component", "Khối giao diện có thể tái sử dụng"],
    ["Server Action", "Hàm chạy trên máy chủ được form hoặc component gọi trực tiếp"],
    ["Middleware", "Mã chạy trước route để kiểm tra hoặc chuyển hướng yêu cầu"],
    ["ORM", "Công cụ ánh xạ bảng database thành đối tượng trong code; Prisma là ORM của dự án"],
    ["Migration", "File SQL có thứ tự để thay đổi cấu trúc database"],
    ["Seed", "Script tạo dữ liệu khởi đầu cần thiết"],
    ["Session", "Bản ghi chứng minh người dùng đã đăng nhập trong một khoảng thời gian"],
    ["Cookie", "Giá trị nhỏ trình duyệt lưu và tự gửi lại cho website"],
    ["Hash", "Kết quả một chiều dùng để tránh lưu mật khẩu hoặc token dạng gốc"],
    ["CRUD", "Create, Read, Update, Delete: tạo, đọc, sửa và xóa dữ liệu"],
    ["Responsive", "Giao diện tự điều chỉnh theo kích thước màn hình"],
], [1.45, 5.05])

doc.add_heading("15 Bảng tra cứu từng file", level=1)
file_rows = [
    ["app/layout.tsx", "Metadata và khung HTML gốc"],
    ["app/globals.css", "Toàn bộ style và breakpoint responsive"],
    ["app/page.tsx", "Trang chủ"],
    ["app/dang-nhap/page.tsx", "Bố cục màn hình đăng nhập"],
    ["app/dang-nhap/login-form.tsx", "Form client và trạng thái gửi"],
    ["app/dang-nhap/actions.ts", "Kiểm tra dữ liệu, mật khẩu và tạo session"],
    ["app/quan-tri/page.tsx", "Dashboard quản trị và requireAdmin"],
    ["app/quan-tri/[section]/page.tsx", "Trang quản trị động theo section"],
    ["components/app-shell.tsx", "Sidebar, topbar và điều hướng chung"],
    ["components/page-head.tsx", "Tiêu đề chuẩn cho trang"],
    ["components/empty-state.tsx", "Trạng thái chưa có dữ liệu"],
    ["lib/prisma.ts", "Một PrismaClient dùng lại trong môi trường dev"],
    ["lib/auth.ts", "Tạo, đọc, xóa session và kiểm tra admin"],
    ["lib/utils.ts", "Chữ viết tắt, định dạng ngày và bỏ dấu"],
    ["middleware.ts", "Chuyển người chưa có cookie về đăng nhập"],
    ["prisma/schema.prisma", "Model, enum, quan hệ và index"],
    ["prisma/seed.ts", "Khởi tạo quản trị viên đầu tiên"],
    ["prisma/migrations/.../migration.sql", "SQL tạo cấu trúc database ban đầu"],
    ["package.json", "Dependency và lệnh npm"],
    ["next.config.ts", "Giới hạn body của Server Action"],
    ["tsconfig.json", "Quy tắc TypeScript và alias @"],
    ["eslint.config.mjs", "Quy tắc lint và thư mục bỏ qua"],
    [".env.example", "Mẫu biến môi trường không chứa bí mật thật"],
    ["SETUP.md", "Hướng dẫn cài đặt và vận hành ngắn"],
]
add_table(doc, ["File", "Trách nhiệm"], file_rows, [2.75, 3.75])

doc.add_heading("Kết luận", level=1)
doc.add_paragraph("Cách học dự án hiệu quả nhất là đi theo một luồng cụ thể. Hãy bắt đầu từ URL, tìm page.tsx, theo component đến Server Action, rồi theo Prisma Client đến model trong schema. Khi thêm chức năng, luôn kiểm tra đầu vào ở máy chủ, kiểm tra quyền trước truy vấn, dùng migration cho thay đổi database và chỉ báo thành công sau khi dữ liệu đã được lưu.")

# Ensure all table headers repeat and avoid row splitting where possible.
for table in doc.tables:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row in table.rows:
        row_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_pr.append(cant_split)

doc.core_properties.title = "Hướng dẫn mã nguồn website DA26TTNT"
doc.core_properties.subject = "Kiến trúc, cách hoạt động và hướng dẫn phát triển cho người mới"
doc.core_properties.author = "Nhóm phát triển DA26TTNT"
doc.save(OUTPUT)
print("Document saved successfully.")
